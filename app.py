
import re
from io import BytesIO

import docx
import fitz
import pandas as pd
import streamlit as st

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# =====================================================
# 1. PAGE CONFIGURATION
# =====================================================

st.set_page_config(
    page_title="AI Resume Screening System",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .block-container {
        padding-top: 1.8rem;
        padding-bottom: 2.5rem;
        max-width: 1250px;
    }

    .hero {
        padding: 1.5rem;
        border-radius: 16px;
        background: linear-gradient(
            120deg, #173b63, #276b83
        );
        color: white;
        margin-bottom: 1.3rem;
    }

    .hero h1 {
        color: white;
        margin-bottom: 0.4rem;
    }

    .hero p {
        color: #e6f2ff;
        margin-bottom: 0;
    }

    div[data-testid="stMetric"] {
        border: 1px solid rgba(128,128,128,0.25);
        padding: 1rem;
        border-radius: 12px;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# =====================================================
# 2. SKILL DICTIONARY
# =====================================================

# Each dictionary key is a standard skill name.
# The values are alternative ways to write that skill.

SKILL_ALIASES = {
    "python": ["python"],
    "java": ["java"],
    "c++": ["c++", "cpp"],
    "c": ["c programming", "c language"],
    "sql": ["sql"],
    "mysql": ["mysql"],
    "machine learning": ["machine learning", "ml"],
    "deep learning": ["deep learning"],
    "artificial intelligence": ["artificial intelligence", "ai"],
    "natural language processing": ["natural language processing", "nlp"],
    "data analysis": ["data analysis", "data analytics"],
    "data science": ["data science"],
    "computer vision": ["computer vision"],
    "html": ["html"],
    "css": ["css"],
    "javascript": ["javascript", "js"],
    "react": ["react", "react.js"],
    "node.js": ["node.js", "nodejs"],
    "flask": ["flask"],
    "django": ["django"],
    "streamlit": ["streamlit"],
    "scikit-learn": ["scikit-learn", "sklearn"],
    "tensorflow": ["tensorflow"],
    "pytorch": ["pytorch", "torch"],
    "pandas": ["pandas"],
    "numpy": ["numpy"],
    "excel": ["excel", "microsoft excel"],
    "power bi": ["power bi", "powerbi"],
    "tableau": ["tableau"],
    "git": ["git"],
    "github": ["github"],
    "teamwork": ["teamwork", "team player"],
    "communication": ["communication", "communication skills"],
    "problem solving": ["problem solving", "problem-solving"],
    "project management": ["project management"],
    "cloud computing": ["cloud computing"],
    "aws": ["aws", "amazon web services"],
    "data visualization": ["data visualization", "data visualisation"],
    "statistics": ["statistics", "statistical analysis"],
    "mongodb": ["mongodb", "mongo db"],
    "postgresql": ["postgresql", "postgres"],
    "rest api": ["rest api", "restful api"],
    "data structures": ["data structures", "data structure"],
    "algorithms": ["algorithms", "algorithm"],
}

EDUCATION_KEYWORDS = [
    "b.e.", "b.e", "b.tech", "bachelor of engineering",
    "bachelor of technology", "b.sc", "bca", "mca",
    "m.tech", "m.e.", "master of science",
    "master of technology", "diploma", "ph.d", "phd",
    "12th", "higher secondary", "10th", "sslc", "puc",
]


# =====================================================
# 3. TEXT EXTRACTION
# =====================================================

def extract_resume_text(uploaded_file):
    """Extract readable text from PDF or DOCX files."""

    filename = uploaded_file.name.lower()

    try:
        if filename.endswith(".pdf"):
            with fitz.open(
                stream=uploaded_file.getvalue(),
                filetype="pdf",
            ) as pdf:
                text_parts = []

                for page in pdf:
                    text_parts.append(page.get_text())

                return "\n".join(text_parts).strip()

        if filename.endswith(".docx"):
            document = docx.Document(
                BytesIO(uploaded_file.getvalue())
            )

            text_parts = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            # Read text in Word tables too.
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_parts.append(cell.text)

            return "\n".join(text_parts).strip()

        return ""

    except Exception as error:
        raise ValueError(
            f"Could not read {uploaded_file.name}: {error}"
        ) from error


# =====================================================
# 4. SKILL EXTRACTION
# =====================================================

def contains_phrase(text, phrase):
    """
    Match a phrase without matching it as part of
    an unrelated word.
    """

    pattern = (
        r"(?<!\w)"
        + re.escape(phrase)
        + r"(?!\w)"
    )

    return re.search(pattern, text, flags=re.IGNORECASE) is not None


def extract_skills(text):
    """Return standard skill names found in the text."""

    found = set()

    for skill, aliases in SKILL_ALIASES.items():
        for alias in aliases:
            if contains_phrase(text, alias):
                found.add(skill)
                break

    return sorted(found)


# =====================================================
# 5. CANDIDATE INFORMATION EXTRACTION
# =====================================================

def extract_email(text):
    match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text,
    )

    return match.group(0) if match else "Not found"


def extract_phone(text):
    patterns = [
        r"(?<!\w)(?:\+91[\s-]?)?[6-9]\d{9}(?!\w)",
        r"(?<!\w)\d{3}[\s.-]\d{3}[\s.-]\d{4}(?!\w)",
    ]

    for pattern in patterns:
        match = re.search(pattern, text)

        if match:
            return match.group(0)

    return "Not found"


def extract_candidate_name(text, filename):
    """
    Heuristic only: the first lines of a resume do not
    always contain the candidate's actual name.
    """

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    excluded = {
        "resume", "curriculum", "vitae",
        "email", "phone", "address",
    }

    for line in lines[:5]:
        cleaned = re.sub(r"\s+", " ", line).strip(" |,-")

        if (
            len(cleaned) <= 60
            and re.fullmatch(r"[A-Za-z][A-Za-z .'-]*", cleaned)
            and len(cleaned.split()) >= 2
            and not any(
                word.lower() in excluded
                for word in cleaned.split()
            )
        ):
            return cleaned

    return filename.rsplit(".", 1)[0]


def extract_education(text):
    lower_text = text.lower()

    found = [
        keyword
        for keyword in EDUCATION_KEYWORDS
        if keyword in lower_text
    ]

    return ", ".join(sorted(set(found))) or "Not detected"


def extract_experience_indicators(text):
    """Find basic experience-related phrases, not verified years."""

    pattern = (
        r"\b(?:internship|intern|work experience|"
        r"professional experience|experience|fresher|"
        r"entry.level|years? of experience)\b"
    )

    matches = re.findall(pattern, text, flags=re.IGNORECASE)

    return ", ".join(sorted(set(
        match.lower() for match in matches
    ))) or "Not detected"


# =====================================================
# 6. NLP SIMILARITY
# =====================================================

def calculate_nlp_similarity(job_description, resume_text):
    """
    Calculate TF-IDF cosine similarity.

    A score of 0 means no shared terms in the
    vectorized text. A higher score indicates more
    textual overlap, not verified job suitability.
    """

    vectorizer = TfidfVectorizer(
        stop_words="english",
        ngram_range=(1, 2),
    )

    try:
        vectors = vectorizer.fit_transform(
            [job_description, resume_text]
        )

        score = cosine_similarity(
            vectors[0:1],
            vectors[1:2],
        )[0][0]

        return round(float(score) * 100, 2)

    except ValueError:
        # For example, both texts may contain only
        # words removed as stop words.
        return 0.0


# =====================================================
# 7. EXPLAINABLE MATCHING SCORE
# =====================================================

def calculate_scores(
    job_description,
    resume_text,
    skill_weight,
):
    """
    Final score combines:
    - Required-skill coverage
    - TF-IDF cosine similarity

    skill_weight is expressed as a percentage.
    """

    job_skills = extract_skills(job_description)
    resume_skills = extract_skills(resume_text)

    matched_skills = sorted(
        set(job_skills) & set(resume_skills)
    )

    missing_skills = sorted(
        set(job_skills) - set(resume_skills)
    )

    if job_skills:
        skill_coverage = (
            len(matched_skills) / len(job_skills)
        ) * 100
    else:
        skill_coverage = 0.0

    nlp_similarity = calculate_nlp_similarity(
        job_description,
        resume_text,
    )

    nlp_weight = 100 - skill_weight

    final_score = (
        skill_coverage * skill_weight / 100
        + nlp_similarity * nlp_weight / 100
    )

    return {
        "skill_coverage": round(skill_coverage, 2),
        "nlp_similarity": round(nlp_similarity, 2),
        "final_score": round(final_score, 2),
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "resume_skills": resume_skills,
    }


# =====================================================
# 8. HEADER AND SIDEBAR
# =====================================================

st.markdown(
    """
    <div class="hero">
        <h1>AI Resume Screening System</h1>
        <p>
            Analyze resumes, understand skill matches,
            and compare candidates using explainable NLP scores.
        </p>
    </div>
    """,
    unsafe_allow_html=True,
)

with st.sidebar:
    st.header("Screening Settings")

    st.write("Adjust how the final score is calculated.")

    skill_weight = st.slider(
        "Required-skill coverage weight (%)",
        min_value=0,
        max_value=100,
        value=60,
        step=5,
    )

    nlp_weight = 100 - skill_weight

    st.metric("NLP similarity weight", f"{nlp_weight}%")

    st.caption(
        "These weights are configurable design choices. "
        "They should be validated before operational use."
    )

    st.divider()

    st.subheader("Supported Files")
    st.write("PDF and DOCX")

    st.subheader("How ranking works")
    st.write(
        "Candidates are sorted by the combined score. "
        "Scores do not establish job performance or "
        "guarantee candidate suitability."
    )


# =====================================================
# 9. JOB DESCRIPTION AND RESUME UPLOAD
# =====================================================

st.subheader("1. Job Description")

job_description = st.text_area(
    "Enter the job requirements",
    height=210,
    placeholder=(
        "Example: Python Developer Intern with Python, SQL, "
        "Machine Learning, NLP, Pandas and NumPy skills..."
    ),
    key="job_description_input",
)

job_skills = extract_skills(job_description)

if job_description.strip():
    if job_skills:
        st.write("**Skills identified in the job description:**")
        st.write(", ".join(job_skills))
    else:
        st.info(
            "No predefined skills were detected. "
            "Add recognizable skills from the supported list "
            "or extend the skill dictionary."
        )

st.subheader("2. Upload Resumes")

uploaded_files = st.file_uploader(
    "Select one or more candidate resumes",
    type=["pdf", "docx"],
    accept_multiple_files=True,
    help=(
        "Text-based PDFs and DOCX files are supported. "
        "Scanned PDFs may require OCR, which is not included."
    ),
)

if uploaded_files:
    st.caption(
        f"{len(uploaded_files)} file(s) selected."
    )


# =====================================================
# 10. ANALYZE AND RANK
# =====================================================

if st.button(
    "Analyze and Rank Candidates",
    type="primary",
    use_container_width=True,
):

    if not job_description.strip():
        st.error(
            "Please enter a job description before analyzing."
        )

    elif not uploaded_files:
        st.error(
            "Please upload at least one PDF or DOCX resume."
        )

    elif not job_skills:
        st.error(
            "No supported skills were identified in the job "
            "description. Add recognizable job skills first."
        )

    else:
        results = []
        resume_details = {}
        errors = []

        progress = st.progress(0)
        status = st.empty()

        total = len(uploaded_files)

        for index, uploaded_file in enumerate(uploaded_files):

            status.write(
                f"Processing {index + 1} of {total}: "
                f"{uploaded_file.name}"
            )

            try:
                resume_text = extract_resume_text(uploaded_file)

                if not resume_text.strip():
                    errors.append(
                        f"{uploaded_file.name}: no readable text "
                        "was found. The file may be scanned or empty."
                    )

                    progress.progress((index + 1) / total)
                    continue

                scores = calculate_scores(
                    job_description,
                    resume_text,
                    skill_weight,
                )

                candidate_name = extract_candidate_name(
                    resume_text,
                    uploaded_file.name,
                )

                results.append({
                    "Candidate": candidate_name,
                    "Resume File": uploaded_file.name,
                    "Email": extract_email(resume_text),
                    "Phone": extract_phone(resume_text),
                    "Education Keywords": extract_education(
                        resume_text
                    ),
                    "Experience Indicators": (
                        extract_experience_indicators(resume_text)
                    ),
                    "Final Match (%)": scores["final_score"],
                    "Skill Coverage (%)": scores["skill_coverage"],
                    "NLP Similarity (%)": scores["nlp_similarity"],
                    "Skills Matched": ", ".join(
                        scores["matched_skills"]
                    ),
                    "Skills Not Found": ", ".join(
                        scores["missing_skills"]
                    ),
                })

                resume_details[uploaded_file.name] = {
                    "text": resume_text,
                    "skills": scores["resume_skills"],
                    "matched": scores["matched_skills"],
                    "missing": scores["missing_skills"],
                }

            except Exception as error:
                errors.append(
                    f"{uploaded_file.name}: {error}"
                )

            progress.progress((index + 1) / total)

        progress.empty()
        status.empty()

        if errors:
            with st.expander(
                f"Files requiring attention ({len(errors)})",
                expanded=True,
            ):
                for error in errors:
                    st.warning(error)

        if not results:
            st.error(
                "No resumes could be analyzed successfully. "
                "Check the file formats and try again."
            )

            st.session_state.pop("screening_results", None)
            st.session_state.pop("resume_details", None)

        else:
            results_df = pd.DataFrame(results)

            results_df = results_df.sort_values(
                by=["Final Match (%)", "Skill Coverage (%)"],
                ascending=[False, False],
            ).reset_index(drop=True)

            results_df.insert(
                0,
                "Rank",
                range(1, len(results_df) + 1),
            )

            st.session_state["screening_results"] = results_df
            st.session_state["resume_details"] = resume_details
            st.session_state["screening_job"] = job_description
            st.session_state["screening_skill_weight"] = skill_weight

            st.success(
                f"Successfully analyzed {len(results_df)} "
                "resume(s)."
            )


# =====================================================
# 11. RESULTS DASHBOARD
# =====================================================

if "screening_results" in st.session_state:

    results_df = st.session_state["screening_results"]

    st.divider()
    st.header("3. Screening Results")

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Resumes Analyzed",
        len(results_df),
    )

    col2.metric(
        "Highest Combined Score",
        f"{results_df['Final Match (%)'].max():.2f}%",
    )

    col3.metric(
        "Required Skills",
        len(extract_skills(
            st.session_state["screening_job"]
        )),
    )

    st.caption(
        "Rankings are based on the combined score. "
        "Review each component and the original resume "
        "before making any recruitment decision."
    )

    display_columns = [
        "Rank",
        "Candidate",
        "Final Match (%)",
        "Skill Coverage (%)",
        "NLP Similarity (%)",
        "Skills Matched",
    ]

    st.dataframe(
        results_df[display_columns],
        use_container_width=True,
        hide_index=True,
    )

    st.subheader("4. Explain an Individual Score")

    selected_file = st.selectbox(
        "Choose a resume to inspect",
        results_df["Resume File"].tolist(),
    )

    candidate = results_df[
        results_df["Resume File"] == selected_file
    ].iloc[0]

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Combined Score",
        f"{candidate['Final Match (%)']:.2f}%",
    )

    col2.metric(
        "Skill Coverage",
        f"{candidate['Skill Coverage (%)']:.2f}%",
    )

    col3.metric(
        "NLP Similarity",
        f"{candidate['NLP Similarity (%)']:.2f}%",
    )

    weight = st.session_state["screening_skill_weight"]
    nlp_weight = 100 - weight

    skill_contribution = (
        candidate["Skill Coverage (%)"] * weight / 100
    )

    nlp_contribution = (
        candidate["NLP Similarity (%)"] * nlp_weight / 100
    )

    st.markdown("**How the final score was calculated**")

    breakdown = pd.DataFrame({
        "Component": [
            "Required-skill coverage",
            "NLP similarity",
        ],
        "Component Score (%)": [
            candidate["Skill Coverage (%)"],
            candidate["NLP Similarity (%)"],
        ],
        "Weight (%)": [
            weight,
            nlp_weight,
        ],
        "Contribution to Final Score": [
            round(skill_contribution, 2),
            round(nlp_contribution, 2),
        ],
    })

    st.dataframe(
        breakdown,
        use_container_width=True,
        hide_index=True,
    )

    st.write(
        f"**Calculation:** "
        f"({candidate['Skill Coverage (%)']:.2f} × "
        f"{weight}%) + "
        f"({candidate['NLP Similarity (%)']:.2f} × "
        f"{nlp_weight}%) = "
        f"**{candidate['Final Match (%)']:.2f}%**"
    )

    st.markdown("**Skills identified in the job description**")
    st.write(
        ", ".join(extract_skills(
            st.session_state["screening_job"]
        )) or "None detected"
    )

    st.markdown("**Skills found in the selected resume**")
    st.write(
        ", ".join(
            st.session_state["resume_details"][selected_file]["skills"]
        ) or "None detected"
    )

    st.markdown("**Matched skills**")
    st.write(
        candidate["Skills Matched"] or "None detected"
    )

    st.markdown("**Job skills not detected in the resume**")
    st.write(
        candidate["Skills Not Found"] or "None"
    )

    st.markdown("**Extracted candidate information**")
    st.write(f"**Name:** {candidate['Candidate']}")
    st.write(f"**Email:** {candidate['Email']}")
    st.write(f"**Phone:** {candidate['Phone']}")
    st.write(
        f"**Education keywords:** "
        f"{candidate['Education Keywords']}"
    )
    st.write(
        f"**Experience indicators:** "
        f"{candidate['Experience Indicators']}"
    )

    with st.expander("View extracted resume text"):
        st.text(
            st.session_state["resume_details"][selected_file]["text"]
        )

    csv_data = results_df.to_csv(
        index=False
    ).encode("utf-8")

    st.download_button(
        "Download Screening Results (CSV)",
        data=csv_data,
        file_name="resume_screening_results.csv",
        mime="text/csv",
        use_container_width=True,
    )


# =====================================================
# 12. FOOTER
# =====================================================

st.divider()

st.caption(
    "AI Resume Screening System | Python | Streamlit | "
    "Scikit-learn | TF-IDF | Cosine Similarity"
)

st.caption(
    "Decision-support prototype only. Matching scores "
    "are not validated predictions of job performance. "
    "Human review is required."
)
